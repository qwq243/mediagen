#!/usr/bin/env python3
"""Versioned image-PPT workspace, approval gates, provenance and delivery.

No inference in init/check/compile/confirm/audit/export/library commands.
Only render --execute --ack-cost may call the existing image engine.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import importlib.util
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import uuid
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
LIBRARY = Path.home() / ".mediagen" / "library"
PACK = "prompt-pack.json"
HEX = re.compile(r"^#[0-9a-fA-F]{6}$")
PHASES = ("content", "prompts", "visual", "deck", "notes")
HUMANS = ("none", "illustrated", "real_photos", "generated_realistic")
PALETTE = {
    "background": "#F5F8FF", "secondary_bg": "#FFFFFF", "primary": "#234EA5",
    "accent": "#14B8A6", "text": "#172B4D", "muted": "#607087", "border": "#DCE5F1",
}


def now():
    return datetime.now(timezone.utc).isoformat()


def sha(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":")).encode()).hexdigest()


def file_sha(path):
    from ppt_read_cache import file_sha as read_sha
    return read_sha(path)


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write(path, value):
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_name(p.name + "." + uuid.uuid4().hex + ".tmp")
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(tmp, p)


def safe_id(value):
    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,80}", value):
        raise ValueError("invalid page/reference id")
    return value


def local(root, name):
    p = Path(name)
    return p.resolve() if p.is_absolute() else (Path(root) / p).resolve()


def asset(root, name):
    p = local(root, name)
    return {"path": str(p), "sha256": file_sha(p)}


def upgrade(pack):
    """Non-destructive normalization. Legacy assertions never become approvals."""
    p = copy.deepcopy(pack)
    if p.get("schema_version", 1) != 3:
        p.setdefault("legacy_schema_version", p.get("schema_version", 1))
        p["schema_version"] = 3
        # Keep old receipts for traceability. New subject definitions never inherit approvals.
        p.setdefault("prior_receipts", {"approvals": copy.deepcopy(p.get("approvals", {})),
                                        "audits": copy.deepcopy(p.get("audits", {}))})
        p["approvals"], p["audits"] = {}, {}
    p.setdefault("revision", 0)
    for key in ("resource_cart", "asset_selection", "revision_requests", "annotations", "composition_records",
                "asset_exports", "edit_plans"):
        p.setdefault(key, [])
    p.setdefault("component_pack", {"enabled": False, "variants": {}, "chapters": [],
                                    "progress_mode": "chapter"})
    p.setdefault("production_route", "full_image")
    p.setdefault("project_id", uuid.uuid4().hex)
    p.setdefault("communication", {
        "audience": "", "communication_intent": "", "audience_outcome": "",
        "core_message": "", "delivery_context": "", "artifact_afterlife": "",
        "duration_seconds": None, "reading_mode": "balanced",
        "page_count": {"mode": "recommend", "count": None, "min": None, "max": None},
        "organization": "mixed", "organization_rationale": "",
    })
    p.setdefault("design", {
        "visual_style": "swiss-minimal", "palette": PALETTE.copy(),
        "typography": {"display": "Microsoft YaHei", "heading": "Microsoft YaHei",
                       "body": "Microsoft YaHei", "numbers": "Arial",
                       "sizes": {"display": 64, "title": 42, "subtitle": 32, "body": 24, "annotation": 16}},
        "icons": "outline", "material": "restrained flat editorial",
        "grid": "12 columns, generous safe margins", "human_policy": "none",
        "prompt_language": "en", "text_in_image": True,
    })
    p["design"].setdefault("baseline_canvas", "1280x720")
    p.setdefault("notes_policy", {"enabled": True, "mode": "generate", "audio": False})
    p.setdefault("transition", {"enabled": False, "effect": "fade", "duration": 0.4,
                                "advance": "manual"})
    p.setdefault("references", [])
    p.setdefault("approvals", {})
    p.setdefault("audits", {})
    p.setdefault("samples", {})
    p.setdefault("generation_authorization", {})
    p.setdefault("generation_calls", [])
    p.setdefault("feedback_status", "pending")
    p.setdefault("primary_language", "zh-CN")
    contract = p.setdefault("generation_contract", {})
    for k, v in {"model": "gpt-image-2", "quality": "high", "size": "2048x1152",
                 "output_format": "png", "image_only_pptx": True}.items():
        contract.setdefault(k, v)
    for page in p.setdefault("pages", []):
        safe_id(page["page_id"])
        page.setdefault("logic", {
            "purpose": "", "key_message_or_question": "",
            "evidence": [], "information_relationship": "",
            "visual_expression": page.get("layout_focus", ""),
            "transition_in": "", "transition_out": "", "duration_seconds": None,
        })
        page.setdefault("speaker_notes", "")
        page.setdefault("materials", [])
        page.setdefault("selected_output", None)
        page.setdefault("output_candidates", [])
        if page.get("transition") is None:
            page.pop("transition", None)
    return p


def load(root):
    raw = (Path(root) / PACK).read_bytes()
    p = upgrade(json.loads(raw.decode("utf-8-sig")))
    p["_disk_sha256"] = hashlib.sha256(raw).hexdigest()
    return p


def save(root, p):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    lock = root / ".mediagen-write.lock"
    try:
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError:
        raise ValueError("workspace is being written; reload after the active writer finishes")
    try:
        target = root / PACK
        if (target.exists() and p.get("_disk_sha256") is not None and
                p["_disk_sha256"] != file_sha(target)):
            raise ValueError("workspace changed since loading; refusing to overwrite another edit")
        if target.exists():
            old = read(target)
            # Immutable, complete snapshots. Never restore call budgets or approvals from these.
            history = root / "revisions" / ("r%06d-" % old.get("revision", 0) + file_sha(target)[:12] + ".json")
            if not history.exists():
                write(history, old)
            p["revision"] = old.get("revision", 0) + 1
        write(target, {k: v for k, v in p.items() if k != "_disk_sha256"})
        p["_disk_sha256"] = file_sha(target)
    finally:
        os.close(fd)
        lock.unlink()


def reference_snapshot(root, refs):
    result = []
    for ref in refs:
        r = copy.deepcopy(ref)
        if r.get("path"):
            r["current_file"] = asset(root, r["path"])
        result.append(r)
    return result


def content_subject(root, p):
    return {
        "communication": p["communication"],
        "primary_language": p["primary_language"],
        "requirements": p.get("requirement_alignment", {}),
        "pages": [{k: v for k, v in page.items() if k in
                   ("page_id", "page_role", "page_text", "logic", "materials",
                    "evidence_boundary", "source_reference", "source_content_injection", "reserved_regions")}
                  for page in p["pages"]],
        "reference_files": reference_snapshot(root, [r for r in p["references"] if r.get("role") == "evidence"]),
    }


def design_hash(root, p):
    return sha({"content": content_subject(root, p), "design": p["design"],
                "generation": p["generation_contract"], "references": reference_snapshot(root, p["references"]),
                "notes_policy": p["notes_policy"],
                "transition": p["transition"],
                "components": p.get("component_pack"), "route": p.get("production_route"),
                "assets": selected_assets(root, p),
                "page_design": [{k: page.get(k) for k in
                                 ("page_id", "layout_focus", "forbidden_items", "reference_ids",
                                    "visual_instructions", "transition", "style_reference_page_id",
                                    "must_preserve", "required_elements", "forbidden_elements",
                                    "reserved_regions")} for page in p["pages"]]})


def selected_assets(root, p):
    return [{**x, "current_sha256": file_sha(local(root, x["path"]))}
            for x in p.get("asset_selection", [])]


def component_values(p, page):
    """Resolve dynamic slots from the current page order; never copy reference-page text."""
    pack = p.get("component_pack", {})
    if not pack.get("enabled"):
        return {}
    variants = pack.get("variants", {})
    variant = variants.get(page.get("page_role", "content"), variants.get("content", {}))
    index = next(i for i, pg in enumerate(p["pages"]) if pg["page_id"] == page["page_id"])
    chapters = pack.get("chapters", [])
    chapter = page.get("chapter", "")
    from ppt_components import applicable
    return {"variant": variant, "items": applicable(pack, page),
            "title": page.get("title") or page.get("page_text", "").split("\n")[0],
            "chapter": chapter, "chapters": chapters, "page_number": index + 1,
            "total_pages": len(p["pages"]), "progress_mode": pack.get("progress_mode", "chapter"),
            "progress": ((chapters.index(chapter) + 1) / len(chapters)
                         if chapters and chapter in chapters and pack.get("progress_mode") == "chapter"
                         else (index + 1) / len(p["pages"]))}


def uses_asset(p, page, item):
    from ppt_components import applicable
    return (not page.get("asset_ids") or item["id"] in page["asset_ids"] or
            item["id"] in {x["asset_id"] for x in applicable(p.get("component_pack", {}), page)})


def page_design_hash(root, p, page):
    """Image validity is page-local; notes, transition and unrelated pages do not repaint it."""
    refs = {x["id"]: x for x in p["references"]}
    used = set(page.get("reference_ids", [])) | {x["reference_id"] for x in page.get("materials", [])}
    anchor = None
    representatives = sample_ids(p)
    if representatives and page["page_id"] not in representatives:
        default_id = representatives[0] if page.get("page_role") in ("cover", "divider", "closing") else representatives[-1]
        anchor_id = page.get("style_reference_page_id", default_id)
        master = next((pg for pg in p["pages"] if pg["page_id"] == anchor_id), {})
        if master.get("selected_output"):
            anchor = asset(root, master["selected_output"]["path"])
    local_design = p["design"]
    if p["design"].get("atomic_spec"):
        from ppt_design_spec import compile_spec
        local_design = {**p["design"], "atomic_spec": compile_spec(p["design"]["atomic_spec"], page["page_id"])}
    return sha({"design": local_design, "generation": p["generation_contract"],
                "reading_mode": p["communication"].get("reading_mode"),
                "page": {k: page.get(k) for k in (
                    "page_id", "page_role", "page_text", "layout_focus", "visual_instructions",
                    "required_elements", "must_preserve", "forbidden_elements", "forbidden_items",
                    "materials", "reference_ids", "style_reference_page_id", "asset_ids", "reserved_regions")},
                "drawing_logic": {k: page.get("logic", {}).get(k) for k in
                                  ("information_relationship", "visual_expression")},
                "components": component_values(p, page), "style_anchor": anchor,
                "references": reference_snapshot(root, [refs[x] for x in sorted(used) if x in refs]),
                "assets": [x for x in selected_assets(root, p) if x["usage"] != "structure"
                           and uses_asset(p, page, x)]})


def prompts(root, p):
    from ppt_coding import prompt_constraints
    result = {}
    by_id = {r["id"]: r for r in p["references"]}
    d = p["design"]
    # Once explicitly activated, quantitative targets have one source of truth.
    # Keep narrative art direction, not a second palette/font/grid instruction set.
    design_keys = ("visual_style", "visual_style_behavior", "icons", "material") if d.get("atomic_spec") else (
        "visual_style", "visual_style_behavior", "palette", "typography", "icons", "material", "grid", "baseline_canvas")
    human = {
        "none": "No people, faces, human silhouettes, photorealistic or illustrated humans.",
        "illustrated": "Illustrated human characters only. No photorealistic people.",
        "real_photos": "Use only supplied authorized photographs. Never invent or redraw a person's identity.",
        "generated_realistic": "Photorealistic generated adults permitted; never impersonate a real person.",
    }[d["human_policy"]]
    for page in p["pages"]:
        refs = []
        components = component_values(p, page)
        if components.get("items"):
            # Source titles/old claims and workflow metadata are not production content.
            components["items"] = [{k: x.get(k) for k in
                ("asset_id", "slot", "scope", "mode", "box", "fixed_appearance_only")}
                for x in components["items"]]
        if not d.get("text_in_image"):
            # Composition owns text slots; do not send competing title/page-number text.
            components = {k: v for k, v in components.items()
                          if k not in ("title", "chapter", "chapters", "page_number", "total_pages")}
        for rid in page.get("reference_ids", []):
            r = by_id[rid]
            refs.append({k: r.get(k) for k in ("id", "role", "use", "preserve", "redraw_allowed")})
        result[page["page_id"]] = "\n".join([
            "Create one professionally art-directed presentation slide as a full-page raster image.",
            f"Canvas: {p['generation_contract']['size']}. Reading context: " +
            p["communication"]["reading_mode"] + ". Respect readable hierarchy; do not shrink text to fit.",
            "DESIGN CONTRACT (appearance guidance, not a guarantee of exact raster font metrics):",
            json.dumps({k: d.get(k) for k in design_keys}, ensure_ascii=False, sort_keys=True),
            "Typography sizes refer to the baseline_canvas; scale proportionally to the output canvas. "
            "Keep identical text roles in the same family/weight across slides; display art lettering "
            "is only for approved cover/divider roles, never a random body-font substitution.",
            human,
            "Maintain the approved style-reference typography, palette, icon/material grammar and margins. "
            "Do not copy its words or blindly duplicate its layout.",
            "PAGE MEANING AND READING ORDER:",
            json.dumps({k: page["logic"].get(k) for k in
                        ("information_relationship", "visual_expression")}, ensure_ascii=False),
            "LAYOUT / EMPHASIS: " + page.get("layout_focus", ""),
            "ADDITIONAL VISUAL INSTRUCTIONS: " + page.get("visual_instructions", ""),
            "REFERENCE ROLES: " + json.dumps(refs, ensure_ascii=False),
            "REVIEWED REFERENCE DESIGN CODES (appearance/structure only; final approved design contract "
            "takes precedence; never render these codes as slide text): " +
            json.dumps(prompt_constraints([x for x in p.get("asset_selection", [])
                if uses_asset(p, page, x)]), ensure_ascii=False),
            "FIXED COMPONENTS AND CURRENT PAGE SLOTS: " + json.dumps(components, ensure_ascii=False),
            "Component references carry appearance only. Never copy their original title, page number, "
            "chapter name or brand into the slide. Use only current approved page slots. Components marked "
            "'composite' will be pasted afterwards: keep their placement areas clear; do not recreate them.",
            "RESERVED LAYOUT ZONES ([x,y,width,height] normalized): " +
            json.dumps([{k: r[k] for k in ("box", "mode", "color", "purpose") if k in r}
                        for r in page.get("reserved_regions", [])], ensure_ascii=False),
            "Keep all slide text, icons and decorative detail outside these zones. They are blank areas "
            "for later insertion, NOT visible placeholders. 'solid' and 'transparent' zones are enforced "
            "after generation. 'image' zones receive original imported image pixels later; do NOT redraw "
            "or invent content inside them. 'prompt_only' zones require visual inspection for compliance.",
            "SUPPLIED ASSET ROLES (use appearance only; never inherit sample claims, numbers or brands): " +
            json.dumps([{"id": x["id"], "kind": x["kind"], "usage": x["usage"]}
                        for x in p.get("asset_selection", []) if x["usage"] != "structure"
                        and x.get("model_allowed") and
                        uses_asset(p, page, x)], ensure_ascii=False),
            "MUST PRESERVE / REQUIRED ELEMENTS: " + json.dumps({
                "must_preserve": page.get("must_preserve", []),
                "required_elements": page.get("required_elements", [])}, ensure_ascii=False),
            "EVIDENCE ZONES: " + json.dumps(page.get("materials", []), ensure_ascii=False),
            "Evidence zones are reserved neutral areas: do NOT synthesize screenshots, logos, certificates "
            "or measured charts. The original evidence will be composited deterministically later.",
            ("EXACT VISIBLE TEXT — retain its original language, punctuation and numerical values. "
             "No added slogans, translations, placeholder text or speaker notes:"
             if d.get("text_in_image") else
             "TEXT-FREE BASE ONLY: do not render letters, numbers or labels. Leave the designated "
             "text areas empty; exact text will be composited deterministically outside this model."),
            ("<visible_text>\n" + page.get("page_text", "") + "\n</visible_text>"
             if d.get("text_in_image") else ""),
            "FORBIDDEN: " + json.dumps({"items": page.get("forbidden_items", []),
                                      "elements": page.get("forbidden_elements", [])}, ensure_ascii=False),
        ])
    if d.get("atomic_spec"):
        from ppt_design_spec import prompt_section, validate_project_spec
        validate_project_spec(d["atomic_spec"], p["pages"])
        for pid in result:
            result[pid] += prompt_section(d["atomic_spec"], pid)
    return result


def output_subject(root, p, page):
    selected = page.get("selected_output")
    if not selected:
        raise ValueError(f"{page['page_id']}: no selected image")
    return {"design": page_design_hash(root, p, page), "prompt": sha(prompts(root, p)[page["page_id"]]),
            "page_id": page["page_id"], "output": asset(root, selected["path"])}


def page_count_errors(communication, actual=None):
    """Stage1 records intent without requiring an already matching storyboard."""
    if "page_count" not in communication:
        return []  # Legacy omission stays omission; do not invalidate old receipts on read.
    pref = communication["page_count"]
    if not isinstance(pref, dict) or pref.get("mode") not in ("recommend", "exact", "range"):
        return ["communication.page_count must use recommend, exact or range"]
    mode = pref["mode"]
    active = {"recommend": (), "exact": ("count",), "range": ("min", "max")}[mode]
    for key in ("count", "min", "max"):
        value = pref.get(key)
        if key in active and (type(value) is not int or value < 1):
            return [f"communication.page_count.{key} must be a positive integer"]
        if key not in active and value is not None:
            return [f"communication.page_count.{key} conflicts with {mode} mode"]
    if mode == "range" and pref["min"] > pref["max"]:
        return ["communication.page_count range must have min <= max"]
    if actual is not None:
        if mode == "exact" and actual != pref["count"]:
            return [f"page count {actual} does not match requested total {pref['count']} (including cover/closing)"]
        if mode == "range" and not pref["min"] <= actual <= pref["max"]:
            return [f"page count {actual} is outside requested total {pref['min']}–{pref['max']} (including cover/closing)"]
    return []


def subject(root, p, phase, page_id=None):
    if phase == "stage1":
        communication = {
            k: p["communication"].get(k) for k in ("audience", "communication_intent",
                "audience_outcome", "core_message", "delivery_context", "artifact_afterlife",
                "content_divergence", "reading_mode", "duration_seconds")}
        if "page_count" in p["communication"]:
            communication["page_count"] = p["communication"]["page_count"]
        return sha({"contract": "communication-v3", "communication": communication,
            "language": p["primary_language"], "canvas": p["generation_contract"]["size"],
            "reference_direction": p.get("reference_direction", "")})
    if phase == "content":
        return sha(content_subject(root, p))
    if phase == "stage2":
        return design_hash(root, p)
    if phase == "prompts":
        return sha({"design": design_hash(root, p), "prompts": prompts(root, p)})
    if phase == "visual":
        page = next(pg for pg in p["pages"] if pg["page_id"] == page_id)
        return sha(output_subject(root, p, page))
    if phase in ("samples", "deck", "notes"):
        pages = p["pages"]
        if phase == "samples":
            ids = sample_ids(p)
            pages = [pg for pg in pages if pg["page_id"] in ids]
        result = {"outputs": [output_subject(root, p, pg) for pg in pages]}
        if phase in ("deck", "notes"):
            result["delivery"] = {"transition": p["transition"], "notes_policy": p["notes_policy"],
                                   "page_transitions": [pg.get("transition") for pg in pages]}
        if phase == "notes":
            result["notes"] = [(pg["page_id"], pg["speaker_notes"]) for pg in pages]
        return sha(result)
    raise ValueError(f"unknown phase: {phase}")


def sample_ids(p):
    pages = p["pages"]
    cover = next((pg["page_id"] for pg in pages if pg.get("page_role") == "cover"), None)
    if not cover and pages:
        cover = pages[0]["page_id"]
    inner = next((pg["page_id"] for pg in pages if pg["page_id"] != cover
                  and pg.get("page_role") not in ("closing", "divider")), None)
    return [x for x in (cover, inner) if x]


def valid(record, expected):
    return bool(record and record.get("status") == "passed" and
                record.get("subject_hash") == expected and record.get("evidence"))


def deviation_errors(root, p, page_ids):
    """Atomic-spec projects require measured, current-image QA, not a status flag."""
    if not p["design"].get("atomic_spec"):
        return []
    import ppt_design_spec as ds
    errors = []
    for pid in page_ids:
        try:
            pg = next(pg for pg in p["pages"] if pg["page_id"] == pid)
            target = ds.compile_spec(p["design"]["atomic_spec"], pid)
            report = p.get("design_deviation_reports", {}).get(pid, {})
            measured = report.get("measurements", {})
            if (not pg.get("selected_output") or report.get("target_hash") != target["spec_hash"] or
                    report.get("output_sha256") != file_sha(local(root, pg["selected_output"]["path"]))):
                raise ValueError("missing or stale target/output-bound report")
            if not measured.get("reviewer") or not measured.get("evidence"):
                raise ValueError("measurement reviewer and evidence required")
            if report.get("status") != "passed" or ds.audit(target, measured)["status"] != "passed":
                raise ValueError("measurements pending or outside agreed tolerances")
        except (ValueError, KeyError, TypeError, StopIteration, OSError) as exc:
            errors.append(f"{pid}: design deviation audit: {exc}")
    return errors


def check(root, p, target="inspect"):
    errors, warnings = [], []
    from ppt_components import validate as validate_components
    from ppt_regions import validate as validate_regions
    try:
        validate_components(p.get("component_pack", {}), p.get("asset_selection", []))
    except (ValueError, KeyError, TypeError) as exc:
        errors.append("component_pack: " + str(exc))
    for item in p.get("asset_selection", []):
        if item.get("protected_evidence") and item.get("model_allowed"):
            errors.append("protected evidence asset must not be sent to the image model")
        if item.get("path") and item.get("sha256") != file_sha(local(root, item["path"])):
            errors.append("pinned asset changed: " + item["id"])
    if target != "inspect" and any(r.get("status") in ("draft", "submitted")
                                   for r in p.get("revision_requests", [])):
        errors.append("pending revision: align and resolve the draft with the agent before production")
    if p.get("retrospective_only") and target != "inspect":
        errors.append("historical retrospective is read-only; prepare a new authorized production job")
    ids = [pg["page_id"] for pg in p["pages"]]
    if not ids or len(ids) != len(set(ids)):
        errors.append("pages must have unique nonempty page ids")
    for field in ("audience", "communication_intent"):
        if not str(p["communication"].get(field, "")).strip():
            errors.append(f"communication.{field} needs clarification")
    if p["communication"].get("reading_mode") not in ("presentation", "text", "balanced"):
        errors.append("reading_mode must be presentation, text or balanced")
    errors.extend(page_count_errors(p["communication"], len(p["pages"])))
    if p["design"].get("human_policy") not in HUMANS:
        errors.append("invalid human_policy")
    if p["design"].get("text_in_image") is not True and p.get("production_route") == "full_image":
        errors.append("full_image requires text_in_image; choose agent_composite/external for deterministic text composition")
    if p["design"].get("prompt_language") != "en":
        warnings.append("non-English visual prompting is a deliberate override, not a proven quality gain")
    for role in PALETTE:
        if not HEX.fullmatch(str(p["design"].get("palette", {}).get(role, ""))):
            errors.append(f"palette.{role} needs #RRGGBB")
    for role in ("display", "heading", "body", "numbers"):
        if not str(p["design"].get("typography", {}).get(role, "")).strip():
            errors.append(f"typography.{role} is required")
    if p["notes_policy"].get("mode") not in ("generate", "optimize", "preserve"):
        errors.append("invalid notes mode")
    for label, override in [("deck", {})] + [(pg["page_id"], pg.get("transition", {})) for pg in p["pages"]]:
        if not isinstance(override, dict):
            errors.append(f"{label}: transition override must be an object")
            continue
        tr = {**p["transition"], **override}
        if tr.get("advance") != "manual":
            errors.append(f"{label}: v2 delivery supports manual advance only; do not silently enable auto advance")
        if tr.get("effect") not in ("fade", "push", "wipe", "split", "cover"):
            errors.append(f"{label}: unsupported curated transition")
        if not isinstance(tr.get("duration"), (int, float)) or not 0.1 <= tr["duration"] <= 5:
            errors.append(f"{label}: transition duration must be 0.1–5 seconds")
        if tr.get("sound"):
            errors.append(f"{label}: sounds are not enabled in this curated transition lane")
    size = re.fullmatch(r"(\d+)x(\d+)", p["generation_contract"].get("size", ""))
    if not size or min(map(int, size.groups())) < 256:
        errors.append("invalid generation size")
    refs = {r["id"]: r for r in p["references"]}
    if len(refs) != len(p["references"]):
        errors.append("duplicate reference id")
    try:
        reference_snapshot(root, p["references"])
    except ValueError as exc:
        errors.append(str(exc))
    for r in p["references"]:
        if target != "inspect" and r.get("registration") == "pending_agent_inspection":
            errors.append(f"{r['id']}: supplied local material awaits agent inspection/import")
        if r.get("send_to_model") and (not r.get("upload_allowed") or not r.get("path")):
            errors.append(f"{r['id']}: model upload requires a local file and explicit permission")
        if r.get("role") == "evidence" and r.get("send_to_model"):
            errors.append(f"{r['id']}: evidence may not enter the image model")
    for pg in p["pages"]:
        pid = pg["page_id"]
        try:
            validate_regions(pg.get("reserved_regions", []), root)
            if any(r["mode"] == "transparent" for r in pg.get("reserved_regions", [])):
                if p.get("production_route") == "full_image":
                    errors.append(f"{pid}: transparent blank base requires agent_composite/external route and final image review")
        except (ValueError, KeyError, TypeError) as exc:
            errors.append(f"{pid}: reserved_regions: {exc}")
        logic = pg["logic"]
        for field in ("purpose", "key_message_or_question", "information_relationship", "visual_expression"):
            if not str(logic.get(field, "")).strip():
                errors.append(f"{pid}: missing logic.{field}")
        seconds = logic.get("duration_seconds")
        if seconds is not None and (not isinstance(seconds, (int, float)) or
                                    not math.isfinite(seconds) or seconds < 0):
            errors.append(f"{pid}: speaking duration must be nonnegative finite seconds or null")
        if not pg.get("page_text", "").strip() and pg.get("page_role") not in ("breathing", "divider"):
            errors.append(f"{pid}: empty visible text")
        for rid in pg.get("reference_ids", []):
            if rid not in refs:
                errors.append(f"{pid}: unknown reference {rid}")
        visual = (pg.get("layout_focus", "") + " " + pg.get("visual_instructions", "") + " " +
                  str(logic.get("visual_expression", "")) + " " +
                  " ".join(refs[rid].get("use", "") for rid in pg.get("reference_ids", []) if rid in refs)).lower()
        if p["design"]["human_policy"] == "none" and re.search(
                r"\b(?:photorealistic person|realistic people|portrait of|human figure)\b|真人|人物照片", visual):
            errors.append(f"{pid}: visual instructions conflict with no-human policy")
        required = pg.get("required_elements", [])
        forbidden = pg.get("forbidden_elements", [])
        if not isinstance(required, list) or not isinstance(forbidden, list):
            errors.append(f"{pid}: required_elements and forbidden_elements must be lists")
        else:
            overlap = {str(x).strip().casefold() for x in required} & {
                str(x).strip().casefold() for x in forbidden}
            if overlap:
                errors.append(f"{pid}: required and forbidden elements conflict: {sorted(overlap)}")
        negatives = " ".join(str(x) for x in pg.get("forbidden_items", [])).lower()
        # Bounded high-confidence conflicts only. Literal page text is data, not a visual command.
        for pattern, name in [
            (r"通栏色带|全宽色带|full[- ]width (?:color )?band", "full-width color band"),
            (r"深色背景|dark background", "dark background"),
            (r"霓虹|neon", "neon"),
        ]:
            if re.search(pattern, visual) and re.search(pattern, negatives):
                errors.append(f"{pid}: requested visual element is also forbidden: {name}")
        fonts = " ".join(str(x) for k, x in p["design"]["typography"].items() if k != "sizes").lower()
        if not re.search(r"正文|body[- ]only", negatives):
            for banned, families, name in [
                (r"宋体|宋楷|serif fonts", r"stzhongsong|simsun|source han serif|noto serif", "serif/Song font"),
                (r"楷体|宋楷|书法体|calligraphy", r"stkaiti|kaiti|calligraph", "Kai/calligraphic font"),
            ]:
                if re.search(banned, negatives) and re.search(families, fonts):
                    errors.append(f"{pid}: global font prohibition conflicts with approved role font: {name}")
        for item in pg["materials"]:
            r = refs.get(item.get("reference_id"))
            box = item.get("box", [])
            if not r or r.get("role") != "evidence" or not r.get("path"):
                errors.append(f"{pid}: evidence overlay needs a registered evidence reference")
            if len(box) != 4 or not all(isinstance(x, (int, float)) and math.isfinite(x) for x in box):
                errors.append(f"{pid}: overlay box must be [x,y,w,h] normalized coordinates")
            elif min(box) < 0 or min(box[2:]) <= 0 or box[0] + box[2] > 1 or box[1] + box[3] > 1:
                errors.append(f"{pid}: overlay box outside canvas")
    for r in p["references"]:
        if p["design"]["human_policy"] == "none" and r.get("human_content") in ("real", "illustrated"):
            errors.append(f"{r['id']}: people-bearing reference conflicts with no-human policy")
    planned = [pg["logic"].get("duration_seconds") for pg in p["pages"]]
    if all(isinstance(x, (int, float)) for x in planned) and p["communication"].get("duration_seconds"):
        if sum(planned) > p["communication"]["duration_seconds"]:
            errors.append("planned slide speaking time exceeds the PPT time budget")
    elif p["communication"].get("duration_seconds"):
        warnings.append("some speaking times are missing; total-time feasibility remains unverified")
    if target != "inspect" and not errors:
        required = ["stage1", "stage2", "content", "prompts"]
        if target in ("render", "export"):
            required.append("samples")
        if target == "export":
            required += ["deck"]
            if p["notes_policy"]["enabled"]:
                required.append("notes")
        for phase in required:
            try:
                expected = subject(root, p, phase)
                record = p["approvals"].get(phase) if phase.startswith("stage") or phase == "samples" else p["audits"].get(phase)
                if not valid(record, expected):
                    errors.append(f"{phase}: missing, failed or stale confirmation/audit")
            except (ValueError, KeyError, StopIteration) as exc:
                errors.append(f"{phase}: {exc}")
        if target == "export":
            for pg in p["pages"]:
                if (pg.get("selected_output") or {}).get("kind") in ("mock", "reference") and not p.get("test_fixture"):
                    errors.append(f"{pg['page_id']}: mock/reference preview cannot masquerade as a delivered model page")
                try:
                    if not valid(p["audits"].get("visual:" + pg["page_id"]),
                                 subject(root, p, "visual", pg["page_id"])):
                        errors.append(f"{pg['page_id']}: missing/failed/stale visual audit")
                except ValueError as exc:
                    errors.append(str(exc))
                if p["notes_policy"]["enabled"] and not pg["speaker_notes"].strip():
                    errors.append(f"{pg['page_id']}: speaker notes missing")
    if p["design"].get("atomic_spec"):
        try:
            import ppt_design_spec as ds
            ds.validate_project_spec(p["design"]["atomic_spec"], p["pages"])
            for pg in p["pages"]:
                ds.compile_spec(p["design"]["atomic_spec"], pg["page_id"])
        except (ValueError, KeyError, TypeError) as exc:
            errors.append("atomic_spec: " + str(exc))
        if target in ("render", "export"):
            errors.extend(deviation_errors(root, p, [pg["page_id"] for pg in p["pages"]]
                                           if target == "export" else sample_ids(p)))
    hashes = {}
    for phase in ("stage1", "stage2", "content", "prompts", "samples", "deck", "notes"):
        try:
            hashes[phase] = subject(root, p, phase)
        except (ValueError, KeyError, StopIteration):
            hashes[phase] = None
    return {"ok": not errors, "errors": errors, "warnings": warnings, "subject_hashes": hashes,
            "sample_page_ids": sample_ids(p), "inference_performed": False}


def confirm(root, p, stage, evidence, origin="chat", expected=None):
    if stage not in ("stage1", "stage2", "samples") or origin not in ("chat", "web"):
        raise ValueError("invalid confirmation stage/origin")
    if not evidence or not str(evidence).strip():
        raise ValueError("explicit user confirmation evidence required")
    if any(r.get("status") in ("draft", "submitted") for r in p.get("revision_requests", [])):
        raise ValueError("pending revision must be aligned with the agent before confirmation")
    if stage == "stage1":
        # The upstream communication round does not require a finished storyboard.
        missing = [k for k in ("audience", "communication_intent", "delivery_context")
                   if not str(p["communication"].get(k, "")).strip()]
        # Existing callers with only audience/intent remain supported; UI asks for context.
        missing = [k for k in missing if k != "delivery_context"]
        errors = ["clarify " + k for k in missing] + page_count_errors(p["communication"])
        inspection = {"ok": not errors, "errors": errors}
    else:
        inspection = check(root, p, "samples" if stage == "samples" else "inspect")
    if not inspection["ok"]:
        raise ValueError("; ".join(inspection["errors"]))
    value = subject(root, p, stage)
    if expected and expected != value:
        raise ValueError("stale confirmation; reload the changed proposal")
    if stage != "stage1" and not valid(p["approvals"].get("stage1"), subject(root, p, "stage1")):
        raise ValueError("stage1 not confirmed for current content")
    if stage == "samples":
        errors = deviation_errors(root, p, sample_ids(p))
        if errors:
            raise ValueError("; ".join(errors))
        for pg in p["pages"]:
            if (pg["page_id"] in sample_ids(p) and
                    (pg.get("selected_output") or {}).get("kind") in ("mock", "reference") and not p.get("test_fixture")):
                raise ValueError("a mock/reference preview is not a real model sample")
            if pg["page_id"] in sample_ids(p) and not valid(
                    p["audits"].get("visual:" + pg["page_id"]), subject(root, p, "visual", pg["page_id"])):
                raise ValueError("representative images need current visual audits before approval")
    p["approvals"][stage] = {"status": "passed", "subject_hash": value,
                             "evidence": [evidence], "origin": origin, "at": now()}
    save(root, p)
    return p["approvals"][stage]


def audit(root, p, report):
    phase, pid = report.get("phase"), report.get("page_id")
    if phase not in PHASES:
        raise ValueError("unknown audit phase")
    if not report.get("independent") or not report.get("reviewer") or not report.get("evidence"):
        raise ValueError("independent reviewer and concrete evidence required")
    expected = subject(root, p, phase, pid)
    if report.get("subject_hash") != expected:
        raise ValueError("audit subject is stale or unspecified")
    findings = report.get("findings", [])
    if any(f.get("severity") in ("blocker", "major") and f.get("resolution") != "fixed"
           for f in findings) and report.get("status") == "passed":
        raise ValueError("unresolved blocking findings cannot pass")
    if report.get("status") not in ("passed", "failed", "pending"):
        raise ValueError("invalid audit status")
    if phase == "visual" and report.get("status") == "passed":
        errors = deviation_errors(root, p, [pid])
        if errors:
            raise ValueError("; ".join(errors))
    key = "visual:" + pid if phase == "visual" else phase
    p["audits"][key] = {**report, "at": now()}
    save(root, p)
    return p["audits"][key]


def compile_pack(root, p):
    result = check(root, p)
    if not result["ok"]:
        raise ValueError("; ".join(result["errors"]))
    records = []
    for pid, prompt in prompts(root, p).items():
        path = Path(root) / "prompts" / f"{pid}.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(prompt + "\n", encoding="utf-8")
        records.append({"page_id": pid, "prompt": prompt, "prompt_sha256": sha(prompt),
                        "size": p["generation_contract"]["size"], "quality": p["generation_contract"]["quality"],
                        "out": str(Path(root) / "outputs/main" / f"{pid}.png")})
    (Path(root) / "prompts/main-prompts.jsonl").write_text(
        "".join(json.dumps(x, ensure_ascii=False) + "\n" for x in records), encoding="utf-8")
    return {"compiled": len(records), "subject_hash": subject(root, p, "prompts"),
            "inference_performed": False}


def select_output(root, p, pid, path, kind, provenance):
    from PIL import Image
    pg = next(pg for pg in p["pages"] if pg["page_id"] == pid)
    if kind not in ("generated", "legacy", "mock", "reference", "user_import", "agent_composite"):
        raise ValueError("invalid image kind")
    full = local(root, path)
    with Image.open(full) as im:
        im.verify()
    with Image.open(full) as im:
        w, h = im.size
    sw, sh = map(int, p["generation_contract"]["size"].split("x"))
    if abs(w / h - sw / sh) > 0.025:
        raise ValueError("selected image aspect ratio differs from canvas; do not distort it")
    if not provenance:
        raise ValueError("output provenance is required")
    if kind == "generated":
        call = next((c for c in reversed(p["generation_calls"]) if
                     c.get("status") == "completed" and c.get("page_id") == pid and
                     c.get("output_path") == str(full) and c.get("output_sha256") == file_sha(full)), None)
        if not call:
            raise ValueError("generated output must match a completed controlled call; use legacy for imported history")
        provenance = "controlled newapi_image.py call " + call["started_at"] + "; authorization " + call["authorization_id"]
    if kind in ("user_import", "agent_composite"):
        candidate = next((c for c in pg.get("output_candidates", [])
                          if c.get("path") == str(full) and c.get("sha256") == file_sha(full)
                          and c.get("kind") == kind), None)
        if not candidate:
            raise ValueError("imported/composited output needs a registered candidate record")
    pg["selected_output"] = {"path": str(full), "sha256_at_selection": file_sha(full),
                             "kind": kind, "provenance": provenance, "selected_at": now()}
    save(root, p)
    return {"subject_hash": subject(root, p, "visual", pid), "output": pg["selected_output"]}


def image_failure_summary(raw):
    """Return an allowlisted diagnosis, never raw output, URLs or credentials."""
    text = str(raw or "").lower()
    if any(term in text for term in ("timeouterror", "timed out", "curl: (28)")):
        return {"category": "timeout"}
    status = re.search(r"\bhttp(?:\s+error)?[\s:]+([45]\d{2})\b", text)
    if status:
        return {"category": "upstream_http", "http_status": int(status.group(1))}
    if any(term in text for term in ("connectionreseterror", "connection reset", "could not resolve host")):
        return {"category": "connection"}
    return {"category": "unclassified"}


def render(root, p, ids, execute=False, ack_cost=False, lane="auto"):
    if not ids or len(set(ids)) != len(ids):
        raise ValueError("render needs unique nonempty page IDs")
    if lane not in ("auto", "samples", "deck"):
        raise ValueError("invalid render lane")
    approved = check(root, p)["subject_hashes"]["samples"]
    already_approved = approved is not None and valid(p["approvals"].get("samples"), approved)
    is_samples = lane == "samples" or (lane == "auto" and not already_approved and
                                      set(ids).issubset(set(sample_ids(p))))
    if is_samples and not set(ids).issubset(set(sample_ids(p))):
        raise ValueError("samples lane only accepts the designated representative pages")
    result = check(root, p, "samples" if is_samples else "render")
    if not result["ok"]:
        raise ValueError("; ".join(result["errors"]))
    compile_pack(root, p)
    plans = []
    for pid in ids:
        safe_id(pid)
        pg = next(pg for pg in p["pages"] if pg["page_id"] == pid)
        out = Path(root) / "outputs/main" / (pid + "-" + uuid.uuid4().hex[:8] + ".png")
        cmd = [sys.executable, str(HERE / "newapi_image.py")]
        # The same ordered resolver drives the web reference strip and actual model inputs.
        from ppt_design_spec import reference_inputs
        images = [Path(item["path"]) for item in reference_inputs(root, p, pg, is_samples)]
        cmd += ["edit" if images else "generate"]
        for im in images:
            cmd += ["--image", str(im)]
        cmd += ["--model", p["generation_contract"]["model"], "--size", p["generation_contract"]["size"],
                "--quality", p["generation_contract"]["quality"], "--output-format", "png",
                "--response-format", "url", "--stream", "--transport", "auto", "--retries", "0",
                "--prompt-file", str(Path(root) / "prompts" / f"{pid}.txt"), "--out", str(out)]
        plans.append((pid, out, cmd))
    if not execute:
        return {"dry_run": True, "calls": [cmd for _, _, cmd in plans], "inference_performed": False}
    auth = p["generation_authorization"]
    if not ack_cost or auth.get("design_hash") != design_hash(root, p) or not auth.get("user_evidence"):
        raise ValueError("paid rendering needs --ack-cost and current explicit generation authorization")
    used = sum(1 for x in p["generation_calls"] if x.get("authorization_id") == auth.get("id"))
    if used + len(plans) > auth.get("max_calls", 0):
        raise ValueError("generation call budget exhausted; no automatic extra calls")
    for pid, out, cmd in plans:
        from ppt_design_spec import request_package, freeze_package
        frozen_request = freeze_package(root, request_package(root, p, pid, is_samples))
        frozen = json.loads(Path(frozen_request).read_text(encoding="utf-8"))
        frozen_dir = Path(frozen_request).parent
        # Execute immutable archived bytes, not a mutable original after checking its hash.
        image_args = iter(frozen["inputs"])
        for index, arg in enumerate(cmd):
            if arg == "--image":
                cmd[index + 1] = str(frozen_dir / next(image_args)["path"])
            elif arg == "--prompt-file":
                cmd[index + 1] = str(frozen_dir / "prompt.txt")
        record = {"page_id": pid, "authorization_id": auth["id"], "started_at": now(), "status": "started",
                  "design_hash": design_hash(root, p), "output_path": str(out.resolve()),
                  "design_request_package": frozen_request}
        p["generation_calls"].append(record)
        save(root, p)  # Count attempted calls even when engine fails.
        completed = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
        record["status"] = "completed" if completed.returncode == 0 and out.exists() else "failed"
        record["return_code"] = completed.returncode
        if record["status"] == "failed":
            record["failure_summary"] = image_failure_summary(
                (getattr(completed, "stderr", "") or "") + "\n" +
                (getattr(completed, "stdout", "") or ""))
        save(root, p)
        if record["status"] == "failed":
            raise ValueError("image engine failed; stopped without retry/model switch; return code recorded, "
                             "safe diagnosis: " + json.dumps(record["failure_summary"]) +
                             "; raw engine output intentionally not persisted because it may contain credentials")
        metadata = {}
        import image_recipe
        recipe = image_recipe.bundle_to(out, Path(root) / "generation-records")
        if recipe:
            metadata["generation_recipe"] = recipe
            record["generation_recipe"] = recipe
        pg = next(pg for pg in p["pages"] if pg["page_id"] == pid)
        from ppt_components import applicable
        has_layout = (any(r["mode"] != "prompt_only" for r in pg.get("reserved_regions", [])) or
                      any(x["mode"] == "composite" for x in applicable(p["component_pack"], pg)))
        if pg["materials"] or has_layout:
            clean = out.with_name(out.stem + "-clean" + out.suffix)
            shutil.copyfile(out, clean)
            metadata.update({"contains_evidence": bool(pg["materials"]) or
                any(r.get("embedded_asset", {}).get("protected_evidence") for r in pg.get("reserved_regions", [])) or any(
                x.get("protected_evidence") and x["id"] in {i["asset_id"] for i in applicable(p["component_pack"], pg)}
                for x in p["asset_selection"]), "clean_model_image": asset(root, clean)})
        if has_layout:
            from ppt_regions import finish_layout
            metadata["composition"] = finish_layout(root, p, pg, out)
        if pg["materials"]:
            compose_evidence(root, p, pid, out)
        record["output_sha256"] = file_sha(out)
        save(root, p)
        import ppt_studio
        ppt_studio.add_candidate(root, p, pid, out, "generated",
                                 "newapi_image.py controlled call " + record["started_at"], metadata)
    return {"generated": len(plans), "inference_performed": True}


def compose_evidence(root, p, pid, output):
    from PIL import Image
    refs = {r["id"]: r for r in p["references"]}
    pg = next(pg for pg in p["pages"] if pg["page_id"] == pid)
    with Image.open(output) as base:
        canvas = base.convert("RGB")
    for item in pg["materials"]:
        r = refs[item["reference_id"]]
        with Image.open(local(root, r["path"])) as source:
            source = source.convert("RGBA")
        x, y, w, h = item["box"]
        box_w, box_h = int(w * canvas.width), int(h * canvas.height)
        source.thumbnail((box_w, box_h), Image.Resampling.LANCZOS)
        px, py = int(x * canvas.width) + (box_w - source.width) // 2, int(y * canvas.height) + (box_h - source.height) // 2
        canvas.paste(source, (px, py), source)
    canvas.save(output)


def register_delivery_download(p, path, label):
    """Register a real delivery artifact for the existing local download route."""
    path = Path(path).resolve()
    if not path.is_file():
        raise ValueError(f"missing delivery file: {path}")
    if path.suffix.lower() not in {".pptx", ".pdf", ".zip", ".md", ".txt", ".html", ".json"}:
        raise ValueError("unsupported delivery file type")
    export_id = "delivery-" + uuid.uuid4().hex[:16]
    link = {
        "label": str(label), "filename": path.name,
        "download_url": f"/api/bundles/{export_id}",
        "sha256": file_sha(path), "bytes": path.stat().st_size,
    }
    p.setdefault("asset_exports", []).append({
        **link, "id": export_id, "kind": "delivery_file", "path": str(path),
        "at": now(), "page_ids": [], "asset_ids": [],
    })
    return link


def export(root, p, dest):
    result = check(root, p, "export")
    if not result["ok"]:
        raise ValueError("; ".join(result["errors"]))
    from pptx import Presentation
    from pptx.util import Inches
    sys.path.insert(0, str(HERE / "vendor/ppt_master"))
    from pptx_transitions import EnterUpdate, AdvanceUpdate, apply_slide_motion, read_slide_transition
    dest = Path(dest)
    if dest.exists():
        raise ValueError("delivery already exists; choose a new version rather than overwrite")
    prs = Presentation()
    w, h = map(int, p["generation_contract"]["size"].split("x"))
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(13.333333 * h / w)
    summaries, manuscript = [], ["# " + p.get("project_name", "PPT") + " — 分页讲稿", "",
                                 "预计时长不等于实际彩排；默认无配音。", ""]
    for pg in p["pages"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(local(root, pg["selected_output"]["path"])), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
        if p["notes_policy"]["enabled"]:
            slide.notes_slide.notes_text_frame.text = pg["speaker_notes"]
            manuscript += ["## " + pg["page_id"], "",
                           f"预计：{pg['logic'].get('duration_seconds')} 秒；实际彩排："
                           f"{pg.get('rehearsal_seconds', '待记录')}", "", pg["speaker_notes"], ""]
        tr = {**p["transition"], **pg.get("transition", {})}
        if tr["enabled"]:
            apply_slide_motion(slide._element,
                               enter=EnterUpdate(effect=tr["effect"], duration=tr["duration"]),
                               advance=AdvanceUpdate(mode="click"))
        summary = read_slide_transition(slide._element)
        summaries.append({"page_id": pg["page_id"], "effect": summary.effect,
                          "duration_ms": summary.duration_ms, "advance_on_click": summary.advance_on_click,
                          "advance_after_ms": summary.advance_after_ms})
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(dest)
    notes_path = dest.with_suffix(".notes.md")
    if p["notes_policy"]["enabled"]:
        notes_path.write_text("\n".join(manuscript), encoding="utf-8")
    # Reopen ZIP through python-pptx, then verify embedded notes/media/transitions.
    reopened = Presentation(dest)
    if len(reopened.slides) != len(p["pages"]):
        raise ValueError("PPTX page count mismatch")
    roundtrip = []
    for slide, pg, expected in zip(reopened.slides, p["pages"], summaries):
        if p["notes_policy"]["enabled"] and slide.notes_slide.notes_text_frame.text != pg["speaker_notes"]:
            raise ValueError("PPTX speaker notes read-back mismatch")
        actual = read_slide_transition(slide._element)
        record = {"page_id": pg["page_id"], "effect": actual.effect,
                  "duration_ms": actual.duration_ms, "advance_on_click": actual.advance_on_click,
                  "advance_after_ms": actual.advance_after_ms}
        if record != expected:
            raise ValueError("PPTX transition read-back mismatch")
        roundtrip.append(record)
    proof = {"pptx": asset(root, dest), "slides": roundtrip,
             "notes": str(notes_path) if p["notes_policy"]["enabled"] else None,
             "powerpoint_wps_visual_compatibility": "not_tested", "at": now()}
    proof["downloads"] = [register_delivery_download(p, dest, "完整PPTX")]
    if p["notes_policy"]["enabled"]:
        proof["downloads"].append(register_delivery_download(p, notes_path, "分页讲稿"))
    proof["feedback_invitation"] = "请使用后反馈视觉、展示逻辑、文字量、讲稿和转场；暂未使用可记为待反馈。"
    write(dest.with_suffix(".delivery.json"), proof)
    p.setdefault("deliveries", []).append(proof)
    save(root, p)
    index_project(root, p)
    return proof


def reject_secrets(value):
    text = json.dumps(value, ensure_ascii=False)
    if re.search(r"\bsk-[A-Za-z0-9_-]{12,}|Bearer\s+\S+|(?:api[_-]?key|access[_-]?token|cookie|password)\s*[\"']?\s*[:=]\s*[\"']?\S+", text, re.I):
        raise ValueError("potential credential in library record; remove it")


def library_record(kind, payload, library=LIBRARY):
    reject_secrets(payload)
    folder = Path(library)
    folder.mkdir(parents=True, exist_ok=True)
    rid = uuid.uuid4().hex
    with closing(sqlite3.connect(folder / "index.sqlite3", timeout=15)) as db:
        with db:
            db.execute("CREATE TABLE IF NOT EXISTS records(id TEXT PRIMARY KEY, kind TEXT, at TEXT, payload TEXT)")
            db.execute("INSERT INTO records VALUES(?,?,?,?)",
                       (rid, kind, now(), json.dumps(payload, ensure_ascii=False)))
    return rid


def library_search(query="", kind=None, library=LIBRARY):
    dbpath = Path(library) / "index.sqlite3"
    if not dbpath.exists():
        return []
    with closing(sqlite3.connect(dbpath)) as db:
        rows = db.execute("SELECT id,kind,at,payload FROM records WHERE payload LIKE ? "
                          "AND (? IS NULL OR kind=?) ORDER BY rowid DESC LIMIT 100",
                          ("%" + query + "%", kind, kind)).fetchall()
    return [{"id": a, "kind": b, "at": c, "payload": json.loads(d)} for a, b, c, d in rows]


def index_project(root, p):
    return library_record("project", {
        "project_id": p["project_id"], "project_name": p.get("project_name"),
        "workspace": str(Path(root).resolve()), "schema_version": 3,
        "design_hash": design_hash(root, p), "communication": p["communication"],
        "design": p["design"], "references": p["references"],
        "pages": [{"page_id": pg["page_id"], "logic": pg["logic"],
                   "selected_output": pg["selected_output"],
                    "prompt_hash": sha(prompts(root, p)[pg["page_id"]]),
                    "prompt_path": str(Path(root).resolve() / "prompts" / (pg["page_id"] + ".txt")),
                    "notes_hash": sha(pg["speaker_notes"]), "transition": pg.get("transition")}
                   for pg in p["pages"]],
        "audit_states": {k: {"status": v.get("status"), "subject_hash": v.get("subject_hash")}
                         for k, v in p["audits"].items()},
        "feedback_status": p["feedback_status"], "deliveries": p.get("deliveries", []),
    })


def record_feedback(root, p, feedback, library=LIBRARY):
    """Keep explicit scopes and child issues/preferences; do not promote guesses to rules."""
    payload = {**feedback, "project_id": p["project_id"],
               "scope": feedback.get("scope", "this_project"), "promotion": "candidate_not_global_rule"}
    if not payload.get("comment", "").strip() and not any(
            payload.get(k) for k in ("preferences", "issues", "lessons")):
        raise ValueError("empty feedback")
    for field in ("preferences", "issues", "lessons"):
        if not isinstance(feedback.get(field, []), list):
            raise ValueError(f"feedback.{field} must be a list")
    if not isinstance(payload["scope"], str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", payload["scope"]):
        raise ValueError("feedback scope must be a short explicit scenario identifier")
    parent = library_record("feedback", payload, library)
    ids = {"feedback": parent}
    for field, kind in (("preferences", "preference"), ("issues", "issue"), ("lessons", "lesson")):
        ids[field] = []
        for entry in feedback.get(field, []):
            item = entry if isinstance(entry, dict) else {"description": str(entry)}
            ids[field].append(library_record(kind, {
                **item, "project_id": p["project_id"], "feedback_id": parent,
                "scope": item.get("scope", payload["scope"]),
                "source": payload.get("source", "user"), "status": item.get("status", "candidate"),
            }, library))
    p["feedback_status"] = "received"
    save(root, p)
    return ids


def register_reference(root, p, ref):
    """Agent/CLI imports an inspected local asset; web paths remain untrusted requests."""
    safe_id(ref["id"])
    if ref.get("role") not in ("style", "content", "evidence"):
        raise ValueError("reference role must be style, content or evidence")
    if not ref.get("license") or not ref.get("use"):
        raise ValueError("reference needs its permission/license and intended use")
    record = copy.deepcopy(ref)
    if record.get("path"):
        src = local(root, record["path"])
        if re.search(r"(?i)(^|[\\/])(?:\\.env(?:\\..*)?|secrets?[^\\/]*|credentials?[^\\/]*|tokens?[^\\/]*|config\\.json)$",
                     str(src)):
            raise ValueError("sensitive configuration/credential file cannot be a material")
        if not src.is_file():
            raise ValueError("reference file does not exist")
        dest = Path(root).resolve() / "assets" / (record["id"] + "-" + file_sha(src)[:12] + src.suffix.lower())
        dest.parent.mkdir(parents=True, exist_ok=True)
        if src != dest:
            shutil.copyfile(src, dest)
        record["path"] = str(dest)
        record["sha256"] = file_sha(dest)
        record["registration"] = "agent_inspected_local_import"
        if record.get("send_to_model"):
            if not record.get("upload_allowed") or record["role"] == "evidence":
                raise ValueError("explicit image-upload consent required; evidence is never uploadable")
            from PIL import Image
            with Image.open(dest) as im:
                im.verify()
    existing = next((r for r in p["references"] if r["id"] == record["id"]), None)
    if existing and existing.get("role") == "evidence" and record["role"] != "evidence":
        raise ValueError("cannot relabel protected evidence as a model style reference")
    p["references"] = [r for r in p["references"] if r["id"] != record["id"]] + [record]
    save(root, p)
    return record


def recommendations(root, p, stage):
    """Bridge the retained upstream UI contract; no model inference or made-up approvals."""
    typeface = p["design"]["typography"]
    palette = {**p["design"]["palette"], "body_text": p["design"]["palette"]["text"],
               "secondary_accent": p["design"]["palette"].get("secondary_accent", p["design"]["palette"]["muted"])}
    typography = {"name": "当前角色字体", "heading": {"primary": typeface["heading"], "english": typeface["numbers"]},
                  "body": {"primary": typeface["body"], "english": typeface["numbers"]},
                  "body_size": typeface["sizes"]["body"], "sizes": typeface["sizes"]}
    sample_check = check(root, p, "samples")
    sample_errors = list(sample_check["errors"])
    audit_states = {}
    for key, record in p["audits"].items():
        parts = key.split(":", 1)
        try:
            expected = subject(root, p, parts[0], parts[1] if len(parts) > 1 else None)
        except (ValueError, KeyError, StopIteration):
            expected = None
        current = bool(expected and valid(record, expected))
        audit_states[key] = {"stored_status": record.get("status"),
                             "effective_status": "passed" if current else (
                                 "stale" if record.get("status") == "passed" else record.get("status", "pending"))}
    for pid in sample_ids(p):
        try:
            if not valid(p["audits"].get("visual:" + pid), subject(root, p, "visual", pid)):
                sample_errors.append(pid + ": representative visual audit missing/failed/stale")
        except ValueError as exc:
            sample_errors.append(str(exc))
    design_confirmed = valid(p["approvals"].get("stage2"), subject(root, p, "stage2"))
    authored = p.get("design_candidates", [])
    current = [
        {"id": "current", "name": "当前推荐方案", "note": p["communication"].get("organization_rationale", "") or
         "需结合资料确认；这是设计示意，不是生图样张。",
         "visual_style": p["design"]["visual_style"], "mode": "custom",
         "mode_behavior": p["communication"].get("organization", "mixed"), "color": {"palette": palette},
         "typography": typography, "display_font": typeface["display"], "icons": p["design"]["icons"],
         "image_strategy": {"rendering": "custom", "behavior": p["design"]["material"]}},
    ]
    # A revised draft must retain the user's choices after an upstream approval becomes stale.
    # Initial agent-authored candidates remain recommendations until the first actual submission.
    if design_confirmed:
        directions = current
    elif p.get("design_choice_recorded"):
        directions = current + authored
    else:
        directions = authored or current
    directions = copy.deepcopy(directions)
    for direction in directions:
        if direction.get("mode") == "custom":
            direction.setdefault("mode_behavior", p["communication"]["organization"])
        direction["delivery_purpose"] = p["communication"]["reading_mode"]
    result = {
        "stage": "stage" + str(stage), "lang": "zh", "primary_language": p["primary_language"],
        "recommend": {"canvas": {"2048x1152": "ppt169", "1536x1152": "ppt43"}.get(p["generation_contract"]["size"], "custom"),
                      "delivery_purpose": p["communication"]["reading_mode"], "mode": "custom",
                      "visual_style": p["design"]["visual_style"], "icons": p["design"]["icons"],
                      "image_usage": ["ai"], "image_strategy": "custom"},
        "custom_candidates": {
            "mode": {"behavior": p["communication"]["organization"]},
            "image_strategy": {"behavior": p["design"]["material"], "name": "当前材质语言"}},
        "canvas": {"recommend": "ppt169" if p["generation_contract"]["size"] == "2048x1152" else p["generation_contract"]["size"]},
        "page_count": {"value": len(p["pages"]), "reason": "根据已确认逐页逻辑调整，不按模板页数凑页"},
        "delivery_purpose": {"recommend": p["communication"]["reading_mode"]},
        "mode": {"recommend": "custom", "custom": {"behavior": p["communication"]["organization"]}},
        "visual_style": {"recommend": p["design"]["visual_style"]},
        "color": {"candidates": [{"name": "当前色卡", "palette": palette}], "selected": 0},
        "typography": {"candidates": [typography], "selected": 0},
        "icons": {"recommend": p["design"]["icons"]},
        "image_usage": {"recommend": ["ai"]},
        "image_strategy": {"recommend": "custom", "candidates": [
            {"rendering": "custom", "behavior": p["design"]["material"], "name": "当前材质语言"}]},
        "design_directions": {"candidates": directions, "selected": 0},
        "template_options": {"library": {}, "explicit": [], "default_mode": "free_design",
                             "options_sha256": sha(p["references"])},
        "mediagen": {"source_hash": design_hash(root, p), "pack": p,
                     "sample_page_ids": sample_ids(p),
                     "stage2_confirmed": valid(p["approvals"].get("stage2"), subject(root, p, "stage2")),
                      "sample_confirmed": valid(p["approvals"].get("samples"), sample_check["subject_hashes"]["samples"]),
                      "sample_subject_hash": sample_check["subject_hashes"]["samples"],
                      "sample_ready": not sample_errors, "sample_blockers": sample_errors,
                      "audit_states": audit_states,
                     "library": library_search(kind="reference"),
                      "history": (library_search(kind="preference")[:8] +
                                  library_search(kind="issue")[:8] + library_search(kind="lesson")[:8]),
                     "sites": [{"name": "SlidesCarnival", "url": "https://www.slidescarnival.com/"},
                               {"name": "GPT-Image2 图库", "url": "https://gpt-image2.canghe.ai/"}]},
    }
    for k in ("audience", "communication_intent", "audience_outcome", "core_message",
              "delivery_context", "artifact_afterlife", "content_divergence"):
        result[k] = {"value": p["communication"].get(k, "")}
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--job", type=Path)
    sub = parser.add_subparsers(dest="command", required=True)
    c = sub.add_parser("init"); c.add_argument("--name", default="图片型 PPT")
    sub.add_parser("migrate")
    c = sub.add_parser("check"); c.add_argument("--target", choices=("inspect", "samples", "render", "export"), default="inspect")
    sub.add_parser("compile")
    c = sub.add_parser("confirm"); c.add_argument("--stage", required=True, choices=("stage1", "stage2", "samples")); c.add_argument("--evidence", required=True)
    c = sub.add_parser("audit"); c.add_argument("--report", type=Path, required=True)
    c = sub.add_parser("select"); c.add_argument("--page", required=True); c.add_argument("--image", required=True)
    c.add_argument("--kind", required=True, choices=("generated", "legacy", "mock", "reference")); c.add_argument("--provenance", required=True)
    c = sub.add_parser("authorize"); c.add_argument("--calls", type=int, required=True); c.add_argument("--evidence", required=True)
    c = sub.add_parser("render"); c.add_argument("--pages", nargs="+", required=True); c.add_argument("--execute", action="store_true"); c.add_argument("--ack-cost", action="store_true")
    c.add_argument("--lane", choices=("auto", "samples", "deck"), default="auto")
    c = sub.add_parser("reference"); c.add_argument("--file", type=Path, required=True)
    c = sub.add_parser("export"); c.add_argument("--out", type=Path, required=True)
    sub.add_parser("index")
    c = sub.add_parser("feedback"); c.add_argument("--file", type=Path, required=True)
    c = sub.add_parser("library"); c.add_argument("--search", default=""); c.add_argument("--kind"); c.add_argument("--add", type=Path)
    args = parser.parse_args(argv)
    try:
        if args.command == "library":
            result = ({"id": library_record(args.kind or "reference", read(args.add))}
                      if args.add else library_search(args.search, args.kind))
        else:
            if not args.job:
                raise ValueError("--job is required")
            root = args.job.resolve()
            if args.command == "init":
                if (root / PACK).exists():
                    raise ValueError("workspace exists; migrate it without overwriting")
                p = upgrade({"project_name": args.name, "pages": []})
                save(root, p)
                result = {"workspace": str(root), "status": "needs_source_content_and_logic"}
            else:
                p = load(root)
                if args.command == "migrate":
                    original = root / PACK
                    backup = original.with_name("prompt-pack.pre-v3-" + uuid.uuid4().hex[:8] + ".json")
                    backup.write_bytes(original.read_bytes())
                    p["schema_version"] = 3
                    save(root, p)
                    result = {"backup": str(backup), "approvals_preserved_as_pending": not bool(p["approvals"])}
                elif args.command == "check":
                    result = check(root, p, args.target)
                elif args.command == "compile":
                    result = compile_pack(root, p)
                elif args.command == "confirm":
                    result = confirm(root, p, args.stage, args.evidence)
                elif args.command == "audit":
                    result = audit(root, p, read(args.report))
                elif args.command == "select":
                    result = select_output(root, p, args.page, args.image, args.kind, args.provenance)
                elif args.command == "authorize":
                    if args.calls < 1 or not args.evidence.strip():
                        raise ValueError("positive call count and explicit user evidence required")
                    p["generation_authorization"] = {"id": uuid.uuid4().hex, "max_calls": args.calls,
                                                     "design_hash": design_hash(root, p), "user_evidence": args.evidence, "at": now()}
                    save(root, p); result = {"max_calls": args.calls, "inference_performed": False}
                elif args.command == "render":
                    result = render(root, p, args.pages, args.execute, args.ack_cost, args.lane)
                elif args.command == "reference":
                    result = register_reference(root, p, read(args.file))
                elif args.command == "export":
                    result = export(root, p, args.out)
                elif args.command == "index":
                    result = {"id": index_project(root, p)}
                elif args.command == "feedback":
                    feedback = read(args.file)
                    result = record_feedback(root, p, feedback)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 1 if isinstance(result, dict) and result.get("ok") is False else 0
    except (ValueError, KeyError, StopIteration, OSError, json.JSONDecodeError) as exc:
        print(json.dumps({"ok": False, "error": str(exc), "automatic_retry": False}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
